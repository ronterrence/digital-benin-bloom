"""Import the Atlas research workbook into deterministic research and public data files."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from collections import Counter
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT = ROOT / "data" / "raw" / "benin_artefacts_master_register.xlsx"
NORMALIZED = ROOT / "data" / "normalized"
PUBLIC_OUTPUT = NORMALIZED / "atlas_public_records.json"
RESEARCH_OUTPUT = NORMALIZED / "atlas_research_records.json"
REPORT_JSON = NORMALIZED / "atlas_import_report.json"
REPORT_CSV = NORMALIZED / "atlas_import_warnings.csv"
RESEARCH_QUEUE_JSON = NORMALIZED / "atlas_research_queue.json"
OVERRIDES = ROOT / "data" / "curated" / "atlas_overrides.json"

HEADERS = [
    "Batch", "Country", "City/Region", "Institution", "Record Level",
    "Object title", "Object type", "Cultural attribution", "Date / period",
    "Material", "Accession / inventory number", "Acquisition date",
    "Provenance summary", "1897 connection", "Digital Benin status/link",
    "Current ownership/restitution status", "Catalogue/source URL",
    "Image available?", "Confidence level", "Research priority", "Notes",
]

CITY_COORDINATES: dict[str, tuple[float, float]] = {
    "Aberdeen": (57.1497, -2.0943), "Boston": (42.3601, -71.0589),
    "Cambridge": (52.2053, 0.1218), "Cambridge, MA": (42.3736, -71.1097),
    "Chicago": (41.8781, -87.6298), "Cleveland": (41.4993, -81.6944),
    "Dallas": (32.7767, -96.7970), "Detroit": (42.3314, -83.0458),
    "Dresden": (51.0504, 13.7373), "Dresden / Leipzig": (51.2200, 13.0800),
    "Edinburgh": (55.9533, -3.1883), "Hamburg": (53.5511, 9.9937),
    "Köln / Cologne": (50.9375, 6.9603), "Leiden": (52.1601, 4.4970),
    "Leiden / Amsterdam / Rotterdam / Berg en Dal": (52.0907, 5.1214),
    "Leipzig": (51.3397, 12.3731), "Minneapolis": (44.9778, -93.2650),
    "New York": (40.7128, -74.0060), "Oxford": (51.7520, -1.2577),
    "Philadelphia": (39.9526, -75.1652), "Providence": (41.8240, -71.4128),
    "Rotterdam": (51.9244, 4.4777), "Saint Louis": (38.6270, -90.1994),
    "Stockholm": (59.3293, 18.0686), "Stuttgart": (48.7758, 9.1829),
    "Vienna": (48.2082, 16.3738), "Washington, D.C.": (38.9072, -77.0369),
}

INSTITUTION_ALIASES = {
    "metropolitan museum of art": "metropolitan-museum-of-art",
    "pitt rivers museum / university of oxford": "pitt-rivers-museum",
    "museum am rothenbaum — markk": "markk-hamburg",
    "museum für kunst und gewerbe hamburg — mk&g": "mkg-hamburg",
    "rautenstrauch-joest-museum": "rautenstrauch-joest-museum",
}


def text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def slug(value: str) -> str:
    value = value.casefold().replace("&", " and ")
    value = re.sub(r"[^a-z0-9]+", "-", value.encode("ascii", "ignore").decode())
    return value.strip("-")


def institution_id(name: str) -> str:
    return INSTITUTION_ALIASES.get(name.casefold(), slug(name))


def looks_like_url(value: str) -> bool:
    return value.startswith("https://") or value.startswith("http://")


def repair_shifted_row(row: dict[str, str], row_number: int) -> tuple[dict[str, str], list[dict[str, Any]]]:
    warnings: list[dict[str, Any]] = []
    # Two source rows omit the 1897 cell, shifting N:U left. This signature is unambiguous.
    if looks_like_url(row["Current ownership/restitution status"]) and row["Catalogue/source URL"].casefold() in {"yes", "likely"}:
        original = row.copy()
        for destination, source in zip(HEADERS[14:], HEADERS[13:20]):
            row[destination] = original[source]
        row["1897 connection"] = ""
        warnings.append({
            "code": "source_columns_repaired", "severity": "warning",
            "field": "1897 connection", "message": f"Workbook row {row_number} had a deterministic N:U column shift; values were realigned and the missing 1897 field remains blank.",
        })
    return row, warnings


def normalize_confidence(value: str) -> str:
    v = value.casefold()
    if "probable" in v or v == "high":
        return "probable"
    if "confirmed" in v:
        return "confirmed"
    if "exclude" in v:
        return "exclude"
    return "unclear"


def normalize_1897(value: str) -> str:
    v = value.casefold()
    if not v:
        return "unclear"
    if "not 1897" in v or "not linked" in v:
        return "not_1897_related"
    if "direct" in v and "1897" in v or "confirmed" in v and "1897" in v:
        return "confirmed_1897_looted"
    if "probable" in v and "1897" in v:
        return "probable_1897_looted"
    if "1897" in v and any(term in v for term in ("possible", "likely", "context", "requires", "needs")):
        return "possible_1897_link"
    if "1897" in v:
        return "benin_object_1897_unproven"
    return "unclear"


def normalize_digital_benin(value: str) -> tuple[str, str, str]:
    url = next(iter(re.findall(r"https?://\S+", value)), "").rstrip(".,;)")
    match = re.search(r"digitalbenin\.org/(?:catalogue/)?(?:objects?/)?([^/?#\s]+)", url, re.I)
    digital_id = match.group(1) if match else ""
    v = value.casefold()
    if url and "probable" in v:
        return "matched_probable", digital_id, url
    if url:
        return "matched_exact", digital_id, url
    if "not found" in v:
        return "not_found", "", ""
    if "not in scope" in v:
        return "not_in_scope", "", ""
    return "not_checked", "", ""


def normalize_statuses(value: str) -> tuple[str, str, str]:
    v = value.casefold()
    ownership = "status_unclear"
    restitution = "unclear"
    physical = "unclear"
    if "ownership" in v and ("transfer" in v or "returned" in v):
        ownership, restitution = "ownership_transferred_to_nigeria", "ownership_transferred"
    elif "physically returned" in v or "repatriated" in v:
        ownership, restitution, physical = "physically_returned_to_nigeria", "physically_returned", "nigeria"
    elif "loan" in v and ("back" in v or "remain" in v):
        ownership, restitution = "returned_but_on_loan_back", "loaned_back"
    elif "review" in v:
        ownership, restitution = "under_restitution_review", "under_review"
    elif any(term in v for term in ("still held", "museum ownership", "retained")):
        ownership, restitution = "still_owned_by_holding_museum", "not_started"
    if "physically returned" in v or "repatriated" in v:
        physical = "nigeria"
    elif "remain" in v or "on display" in v or "on loan" in v:
        physical = "holding_institution_or_loan"
    return ownership, restitution, physical


def make_warning(record_id: str, code: str, field: str, message: str, severity: str = "warning") -> dict[str, str]:
    return {"record_id": record_id, "code": code, "field": field, "severity": severity, "message": message}


TASK_DEFINITIONS = {
    "digital_benin_unresolved": ("missing_digital_benin_id", "Check Digital Benin match", "high"),
    "unclear_1897_status": ("unclear_1897_status", "Verify 1897 relationship", "high"),
    "unclear_ownership": ("unclear_ownership", "Verify legal ownership and restitution status", "high"),
    "missing_image_rights": ("missing_image_rights", "Document image rights", "medium"),
    "missing_coordinates": ("missing_coordinates", "Add or verify institution coordinates", "medium"),
}


def research_tasks(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    tasks = []
    for record in records:
        for warning in record.get("warnings", []):
            definition = TASK_DEFINITIONS.get(warning["code"])
            if not definition:
                continue
            task_type, title, priority = definition
            tasks.append({
                "id": f"task-{record['local_record_id']}-{warning['code']}",
                "object_id": record["local_record_id"],
                "institution_id": record["institution_id"],
                "task_type": task_type,
                "title": title,
                "description": warning["message"],
                "priority": priority,
                "status": "todo",
                "source_urls": record.get("source_urls", []),
                "created_at": date.today().isoformat(),
                "updated_at": date.today().isoformat(),
            })
    return tasks


def validate(record: dict[str, Any]) -> list[dict[str, str]]:
    rid = record["local_record_id"]
    warnings: list[dict[str, str]] = []
    def add(code: str, field: str, message: str, severity: str = "warning") -> None:
        warnings.append(make_warning(rid, code, field, message, severity))
    if not record["source_urls"]: add("missing_source", "source_urls", "Record has no source URL.", "error")
    if record["record_level"] == "object" and not record["accession_number"]: add("missing_accession", "accession_number", "Object has no accession or inventory number.")
    if not record["institution_name"]: add("missing_institution", "institution_name", "Record has no institution.", "error")
    if record["expedition_1897_status"] == "unclear": add("unclear_1897_status", "expedition_1897_status", "The relationship to the 1897 expedition is unresolved.")
    if record["digital_benin_match_status"] in {"not_checked", "not_found"}: add("digital_benin_unresolved", "digital_benin_match_status", "Digital Benin matching work remains unresolved.")
    if record["digital_benin_match_status"] == "matched_exact" and not record["digital_benin_url"]: add("digital_benin_url_missing", "digital_benin_url", "Exact match has no Digital Benin URL.", "error")
    if record["current_ownership_status"] == "status_unclear": add("unclear_ownership", "current_ownership_status", "Current legal ownership is unclear.")
    if record["current_physical_location_status"] == "unclear": add("unclear_physical_location", "current_physical_location", "Current physical location is unclear.")
    if record["image_available"] and not record["image_rights"]: add("missing_image_rights", "image_rights", "An image is reported available, but rights are not documented.")
    if record["record_level"] == "collection": add("collection_level_record", "record_level", "This row describes a collection or count, not one individual object.", "info")
    culture = record["cultural_attribution_original"].casefold()
    if "republic of benin" in culture or culture == "benin": add("benin_geography_ambiguity", "cultural_attribution_original", "Check Republic of Benin versus Kingdom of Benin attribution.")
    return warnings


def make_record(row: dict[str, str], row_number: int) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    row, repair_warnings = repair_shifted_row(row, row_number)
    level = "collection" if row["Record Level"].casefold() == "collection" else "object"
    inst_id = institution_id(row["Institution"])
    identity = row["Accession / inventory number"] or f'{row["Object title"]}-{row["Object type"]}'
    digest = hashlib.sha1(f"{level}|{inst_id}|{identity.casefold()}".encode()).hexdigest()[:10]
    rid = f"atlas-{inst_id}-{digest}"
    db_status, db_id, db_url = normalize_digital_benin(row["Digital Benin status/link"])
    ownership, restitution, physical = normalize_statuses(row["Current ownership/restitution status"])
    image_text = row["Image available?"].casefold()
    coordinates = CITY_COORDINATES.get(row["City/Region"])
    source_url = row["Catalogue/source URL"] if looks_like_url(row["Catalogue/source URL"]) else ""
    record: dict[str, Any] = {
        "local_record_id": rid, "record_level": level, "review_state": "imported_unreviewed",
        "country": row["Country"], "city": row["City/Region"],
        "institution_name": row["Institution"], "institution_id": inst_id,
        "latitude": coordinates[0] if coordinates else None, "longitude": coordinates[1] if coordinates else None,
        "coordinate_precision": "city" if coordinates else "unknown",
        "coordinate_source": "Curated city centroid; verify before object-level geographic analysis" if coordinates else "",
        "object_title": row["Object title"], "object_type": row["Object type"],
        "object_type_normalized": slug(row["Object type"]).replace("-", " "),
        "cultural_attribution_original": row["Cultural attribution"],
        "cultural_attribution_normalized": "Edo / Kingdom of Benin" if any(x in row["Cultural attribution"].casefold() for x in ("edo", "kingdom of benin", "benin kingdom")) else "Unresolved",
        "date_period": row["Date / period"], "material": row["Material"],
        "accession_number": row["Accession / inventory number"] if level == "object" else "",
        "collection_identifier": row["Accession / inventory number"] if level == "collection" else "",
        "acquisition_date": row["Acquisition date"], "provenance_summary": row["Provenance summary"],
        "provenance_confidence": normalize_confidence(row["Confidence level"]),
        "confidence_level": normalize_confidence(row["Confidence level"]),
        "expedition_1897_status": normalize_1897(row["1897 connection"]),
        "expedition_1897_source_text": row["1897 connection"],
        "digital_benin_match_status": db_status, "digital_benin_id": db_id, "digital_benin_url": db_url,
        "digital_benin_source_text": row["Digital Benin status/link"],
        "current_ownership_status": ownership, "restitution_status": restitution,
        "current_physical_location_status": physical,
        "current_physical_location": "Nigeria" if physical == "nigeria" else row["Institution"] if physical == "holding_institution_or_loan" else "",
        "ownership_restitution_source_text": row["Current ownership/restitution status"],
        "museum_catalogue_url": source_url, "source_urls": [source_url] if source_url else [],
        "image_available": image_text.startswith("yes") or "image" in image_text,
        "image_url": "", "image_rights": "", "image_credit": "",
        "research_priority": row["Research priority"].casefold() if row["Research priority"].casefold() in {"high", "medium", "low"} else "medium",
        "public_notes": "", "source_notes": "",
        "last_verified_date": "", "legacy_references": [],
        "warnings": [], "_import_warnings": [],
        "_internal_research_notes": row["Notes"], "_raw_import_payload": row,
        "_source_row": row_number,
    }
    warnings = [make_warning(rid, w["code"], w["field"], w["message"], w["severity"]) for w in repair_warnings]
    if not coordinates:
        warnings.append(make_warning(rid, "missing_coordinates", "city", "No curated city coordinates are available."))
    warnings.extend(validate(record))
    record["warnings"] = warnings
    record["_import_warnings"] = warnings[:len(repair_warnings) + (1 if not coordinates else 0)]
    return record, warnings


def read_workbook(path: Path) -> list[tuple[dict[str, str], int]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "Master Register" not in workbook.sheetnames:
        raise ValueError("Workbook is missing the 'Master Register' sheet")
    sheet = workbook["Master Register"]
    rows = list(sheet.iter_rows(values_only=True))
    actual = [text(v) for v in rows[0]]
    if actual != HEADERS:
        raise ValueError(f"Unexpected workbook headers: {actual}")
    return [(dict(zip(HEADERS, [text(v) for v in values])), number) for number, values in enumerate(rows[1:], 2) if any(v is not None for v in values)]


def load_json(path: Path, default: Any) -> Any:
    if not path.exists(): return default
    return json.loads(path.read_text(encoding="utf-8"))


def matching_keys(record: dict[str, Any]) -> list[str]:
    keys = []
    if record.get("digital_benin_id"): keys.append("dbid:" + record["digital_benin_id"].casefold())
    if record.get("digital_benin_url"): keys.append("dburl:" + record["digital_benin_url"].casefold())
    inst = record.get("institution_id", "")
    if record.get("accession_number"): keys.append(f"accession:{inst}:{record['accession_number'].casefold()}")
    if record.get("collection_identifier"): keys.append(f"collection:{inst}:{record['collection_identifier'].casefold()}")
    if record.get("museum_catalogue_url"): keys.append("museum:" + record["museum_catalogue_url"].casefold())
    keys.append(f"fallback:{inst}:{slug(record.get('object_title',''))}:{slug(record.get('object_type',''))}")
    return keys


CURATED_FIELDS = {
    "digital_benin_id", "digital_benin_url", "digital_benin_source_text",
    "museum_catalogue_url", "source_urls", "source_notes", "public_notes",
    "image_url", "image_rights", "image_credit", "last_verified_date",
}


def merge_curated_fields(incoming: dict[str, Any], previous: dict[str, Any]) -> dict[str, Any]:
    """Keep prior curated values when the incoming workbook leaves them blank."""
    merged = dict(incoming)
    for field in CURATED_FIELDS:
        old_value = previous.get(field)
        new_value = merged.get(field)
        old_present = bool(old_value) if not isinstance(old_value, list) else bool(old_value)
        new_present = bool(new_value) if not isinstance(new_value, list) else bool(new_value)
        if old_present and not new_present:
            merged[field] = old_value
    return merged


def apply_overrides(records: list[dict[str, Any]]) -> None:
    overrides = load_json(OVERRIDES, {})
    for record in records:
        patch = overrides.get(record["local_record_id"], {})
        for key, value in patch.items():
            if not key.startswith("_"):
                record[key] = value
        record["warnings"] = record.get("_import_warnings", []) + validate(record)


def public_record(record: dict[str, Any]) -> dict[str, Any]:
    return {key: value for key, value in record.items() if not key.startswith("_")}


def import_atlas(input_path: Path, output_dir: Path = NORMALIZED) -> dict[str, Any]:
    output_dir.mkdir(parents=True, exist_ok=True)
    previous = load_json(output_dir / RESEARCH_OUTPUT.name, [])
    previous_index = {key: record for record in previous for key in matching_keys(record)}
    records, warnings, duplicates, changed_fields = [], [], [], 0
    seen: dict[str, str] = {}
    created = updated = skipped = 0
    for row, row_number in read_workbook(input_path):
        record, row_warnings = make_record(row, row_number)
        matched = next((previous_index[key] for key in matching_keys(record) if key in previous_index), None)
        if matched:
            record = merge_curated_fields(record, matched)
            changes = sum(matched.get(k) != v for k, v in record.items() if not k.startswith("_"))
            if changes: updated += 1; changed_fields += changes
            else: skipped += 1
        else: created += 1
        accession_key = f"{record['institution_id']}|{record['accession_number'] or record['collection_identifier']}".casefold()
        if accession_key in seen:
            warning = make_warning(record["local_record_id"], "possible_duplicate", "accession_number", f"Possible duplicate of {seen[accession_key]}.", "error")
            row_warnings.append(warning); record["warnings"].append(warning); record["_import_warnings"].append(warning); duplicates.append([seen[accession_key], record["local_record_id"]])
        else: seen[accession_key] = record["local_record_id"]
        records.append(record); warnings.extend(row_warnings)
    apply_overrides(records)
    warnings = [warning for record in records for warning in record["warnings"]]
    public = [public_record(record) for record in records]
    queue = research_tasks(records)
    (output_dir / RESEARCH_OUTPUT.name).write_text(json.dumps(records, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (output_dir / PUBLIC_OUTPUT.name).write_text(json.dumps(public, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (output_dir / RESEARCH_QUEUE_JSON.name).write_text(json.dumps(queue, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    report = {
        "source_file": input_path.name, "generated_on": date.today().isoformat(),
        "total_records": len(records), "object_records": sum(r["record_level"] == "object" for r in records),
        "collection_records": sum(r["record_level"] == "collection" for r in records),
        "records_created": created, "records_updated": updated, "records_skipped": skipped,
        "possible_duplicates": duplicates, "fields_changed": changed_fields,
        "research_tasks_created": len(queue),
        "warnings_created": len(warnings), "errors": [w for w in warnings if w["severity"] == "error"],
        "warning_counts": dict(sorted(Counter(w["code"] for w in warnings).items())),
    }
    (output_dir / REPORT_JSON.name).write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    with (output_dir / REPORT_CSV.name).open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["record_id", "severity", "code", "field", "message"])
        writer.writeheader(); writer.writerows(warnings)
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", nargs="?", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output-dir", type=Path, default=NORMALIZED)
    args = parser.parse_args()
    report = import_atlas(args.input.resolve(), args.output_dir.resolve())
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
